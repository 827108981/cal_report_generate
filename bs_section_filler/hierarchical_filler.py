from __future__ import annotations

import csv
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Optional

from docx import Document

from .excel_blocks import ExcelBlock, SourceCell, read_excel_blocks
from .utils import (
    compact_texts,
    format_value,
    is_blank_value,
    normalize_text,
    looks_like_top_section,
    physical_table_matrix,
    row_nonblank_positions,
    row_physical_tcs,
    safe_ratio,
    set_tc_text_keep_style,
    set_table_fixed_layout,
    tc_is_writable,
    tc_text,
)
from .word_blocks import WordTableBlock, get_word_table_blocks
from .word_formula_check import find_word_table_formulas


@dataclass
class FillLog:
    word_table_index: int
    word_section: str
    word_object: str
    word_row: int
    word_col: int
    module: str
    excel_sheet: str
    excel_section: str
    excel_object: str
    excel_row: int
    excel_col: int
    value: str
    reason: str


@dataclass
class MatchLog:
    word_table_index: int
    word_section: str
    word_object: str
    module: str
    excel_sheet: str
    excel_start_row: int
    excel_end_row: int
    excel_section: str
    excel_object: str
    score: float
    note: str


def _contains(a: str, b: str) -> bool:
    if not a or not b:
        return False
    return a in b or b in a


def _canonical_row_label(text: str) -> str:
    """Normalize business labels that differ between Excel and Word templates.

    The BS-5000 workbook and Word template use several equivalent names, e.g.
    "参考色素最大值" vs "参考色素均值" and "γ-GT" vs "GGT".  These
    aliases must only affect row-key matching; the displayed Word labels remain
    unchanged.
    """
    value = normalize_text(text)
    if not value:
        return ""
    aliases = {
        "ALTALTN": "ALT",
        "Γ-GT": "GGT",
        "ΓGT": "GGT",
        "γ-GT": "GGT",
        "γGT": "GGT",
        "LDHL->P": "LDH",
        "LDH(L->P)": "LDH",
        "UREAUREAN": "UREA",
        "TPTPN": "TP",
        "TB-D": "TBD",
        "T-BIL-V": "TBD",
    }
    if value in aliases:
        return aliases[value]
    if value.startswith("ALT") and "ALTN" in value:
        return "ALT"
    if value in {"Γ-GT", "ΓGT"} or "γ-GT" in text or "γGT" in text:
        return "GGT"
    if value.startswith("LDH"):
        return "LDH"
    if value.startswith("UREA"):
        return "UREA"
    if value.startswith("TP") and "TPT" in value:
        return "TP"
    # These labels describe the same row but express the threshold differently
    # in the source workbook and the Word form.  Normalizing them here keeps
    # the value offset anchored to the real Excel label instead of a fallback.
    if value.startswith("\u6700\u5927\u5438\u5149\u5ea6"):
        return "MAX_ABSORBANCE_THRESHOLD"
    if "\u52a0\u6837\u8bef\u5dee" in value or "\u52a0\u6837\u51c6\u786e\u5ea6\u6307\u6807" in value:
        return "SAMPLING_ACCURACY_INDICATOR"
    if "参考色素" in value and ("均值" in value or "最大值" in value):
        return "参考色素值"
    if "携带污染率" in value and not any(x in value for x in ("CLH", "CHL", "指标", "不大于", "小于")):
        return "携带污染率值"
    return value


def _block_kind_from_rows(rows: list[list[str]]) -> str:
    flat = [normalize_text(cell) for row in rows for cell in row if cell]
    joined = "|".join(flat)
    first_cells = [normalize_text(row[0]) for row in rows if row and row[0]]
    numeric_rows = sum(1 for value in first_cells if value.isdigit())

    # 加样检测存在两种完全不同的方法：
    # 1) 光度法：吸光度/结果；
    # 2) 称重法：M0、M1、加样量。
    # 两者虽然处于同一章节，但数据不能互相填充。
    if all(token in joined for token in ("M0", "M1", "加样量")):
        return "gravimetric_sampling"
    if "吸光度" in joined and "结果" in joined and ("样本针" in joined or "试剂针" in joined):
        return "photometric_sampling"

    if all(token in joined for token in ("质控靶值", "实测值")) and ("范围低限" in joined or "范围高限" in joined):
        return "clinical_accuracy"
    if numeric_rows >= 10 and all(token in joined for token in ("MEAN", "SD", "CV")):
        return "clinical_precision"
    if "橙黄G样本针携带污染" in joined and "第1组" in joined:
        return "sample_carryover"
    return ""


def _word_block_kind(word: WordTableBlock) -> str:
    return _block_kind_from_rows(word.matrix)


def _excel_block_kind(excel: ExcelBlock) -> str:
    return _block_kind_from_rows([[cell.text for cell in row] for row in excel.rows])


def _looks_like_tail_inventory_word_table(word: WordTableBlock) -> bool:
    """Identify appendix-like reagent/consumable tables that must not fuzzy-match data blocks."""
    head = normalize_text(" ".join(cell for row in word.matrix[:4] for cell in row if cell))
    return any(token in head for token in (
        "当前使用的校准溶液",
        "校准品质控品名称",
        "项目名称批号失效期",
        "生产厂商失效期",
    ))


def _score_block(word: WordTableBlock, excel: ExcelBlock) -> float:
    score = 0.0
    ws = normalize_text(word.section)
    es = normalize_text(excel.section)
    wo = normalize_text(word.object_name)
    eo = normalize_text(excel.object_name)
    wt = normalize_text(word.title)
    et = normalize_text(excel.title)
    section_related = bool(ws and es and _contains(ws, es))
    object_related = bool(wo and eo and _contains(wo, eo))
    title_related = bool(wt and et and _contains(wt, et))

    if ws and es:
        if ws == es:
            score += 60
        elif _contains(ws, es):
            score += 38
    # 对象名是关键，例如 0.5A橙黄G溶液
    if wo and eo:
        if wo == eo:
            score += 70
        elif _contains(wo, eo):
            score += 48
    if wt and et:
        if wt == et:
            score += 22
        elif _contains(wt, et):
            score += 14
    # 表头相似度作为兜底
    header_ratio = safe_ratio(word.header_norms, excel.header_norms)
    score += header_ratio * 30

    # 大型 Word 表不能匹配到只有一行说明文字的 Excel 锚点。
    if len(word.matrix) >= 4 and len(excel.rows) <= 1:
        score -= 100

    # 对临床精密度/准确性、样本携带污染率等结构相近但标题重复的区域，
    # 以表格结构作为强信号，避免仅凭章节名选中错误说明块。
    word_kind = _word_block_kind(word)
    excel_kind = _excel_block_kind(excel)
    if word_kind and excel_kind:
        score += 120 if word_kind == excel_kind else -100

    # 有些尾部清单表只有章节名与前面的检测表相同，本身没有可核对的数据表头。
    # 这种“章节名单点命中”不能跨表借数据，否则会把电解质准确度等数据写进失效期/批号表。
    if (
        section_related
        and not object_related
        and not title_related
        and header_ratio == 0
        and not (word_kind and excel_kind)
        and _looks_like_tail_inventory_word_table(word)
    ):
        score = min(score, 20)

    # 如果完全无章节或对象匹配，只靠表头不允许高分
    if not (section_related or object_related or title_related):
        score = min(score, 20)
    return score


def _choose_block(word: WordTableBlock, blocks: list[ExcelBlock], module: str) -> tuple[Optional[ExcelBlock], float]:
    candidates = [b for b in blocks if b.module == module]
    if not candidates:
        return None, 0.0
    best = None
    best_score = -1.0
    for b in candidates:
        s = _score_block(word, b)
        if s > best_score:
            best = b
            best_score = s
    # 低于阈值不要强行填，避免填错表
    if best is None or best_score < 55:
        return None, best_score
    return best, best_score


def _exact_named_block(
    blocks: list[ExcelBlock],
    module: str,
    section_name: str,
    block_name: str,
) -> Optional[ExcelBlock]:
    """Return one exact Excel sub-block without fuzzy fallback.

    Some Word tables contain two independent result regions in one physical table.
    BS-5000 dark-current detection is one such case: the Word table contains both
    inner-ring and outer-ring regions, while Excel stores them as two separate
    blocks.  Fuzzy selection of a single block inevitably leaves one region empty
    or duplicates the other.  This helper deliberately uses exact section/name
    matching so the local fix cannot affect unrelated sections.
    """
    expected_section = normalize_text(section_name)
    expected_name = normalize_text(block_name)
    for block in blocks:
        if block.module != module:
            continue
        if normalize_text(block.section) != expected_section:
            continue
        names = {normalize_text(block.title), normalize_text(block.object_name)}
        if expected_name in names:
            return block
    return None


def _dark_current_split_sources(
    word: WordTableBlock,
    blocks: list[ExcelBlock],
    module: str,
) -> tuple[Optional[ExcelBlock], Optional[ExcelBlock]] | None:
    """Resolve the two Excel sources for the combined BS-5000 dark-current table.

    The template uses one Word table with two regions (inner/outer), but the Excel
    reader emits two independent blocks.  Handle only this exact shape and leave
    the generic matcher unchanged for every other report section.
    """
    if normalize_text(word.section) != normalize_text("暗电流检测"):
        return None
    joined = normalize_text(" ".join(cell for row in word.matrix for cell in row if cell))
    if normalize_text("内圈暗电流测试") not in joined or normalize_text("外圈暗电流测试") not in joined:
        return None
    inner = _exact_named_block(blocks, module, "暗电流检测", "内圈暗电流测试")
    outer = _exact_named_block(blocks, module, "暗电流检测", "外圈暗电流测试")
    return inner, outer


def _source_row_texts(row: list[SourceCell]) -> list[str]:
    return [c.text for c in row if not is_blank_value(c.value)]


def _looks_like_numeric_header_source_row(row: list[SourceCell]) -> bool:
    """识别 Excel 中类似 1/2/3 的二层表头，避免 Word 行键 1 错匹配到表头。

    v4 用“数字行键必须在 Excel 第一列”来规避该问题，但这会误杀
    吸光度重复性/稳定性等表，因为这些真实数据行的次数在 Excel 第4列。
    这里改成只跳过“整行都是小整数序号”的表头行。
    """
    texts = [normalize_text(c.text) for c in row if not is_blank_value(c.value)]
    if len(texts) < 2:
        return False
    if not all(t.isdigit() for t in texts):
        return False
    nums = [int(t) for t in texts]
    return nums == list(range(nums[0], nums[0] + len(nums))) and max(nums) <= 20


HIGH_LEVEL_ROW_LABELS = {'测试数据', '准确性计算', '重复性计算', '稳定性计算'}


def _is_placeholder_text(t: str) -> bool:
    n = normalize_text(t)
    return n in {'', '/', '／', '是否', 'PASS', 'FAIL', '■PASS', '■FAIL', 'PASSFAIL', '■PASSFAIL', '□PASS□FAIL', '通过不通过', '合格不合格', 'N/A'}

def _row_key_candidates(texts: list[str]) -> list[str]:
    """从 Word 行中抽取行键：优先数字/编号，其次 MAX/MIN/CV/结论等。"""
    clean = [t.strip() for t in texts if t and t.strip() and not _is_placeholder_text(t)]
    if not clean:
        return []
    keys: list[str] = []
    for t in clean:
        n = normalize_text(t)
        if n in {'测试数据', '准确性计算', '重复性计算', '稳定性计算'}:
            continue
        # 数字行键、MAX/MIN/CV/结论等
        if n:
            keys.append(n)
    # 同时返回组合键，便于“稳定性计算 + MAX”匹配
    combo = normalize_text(''.join(clean[:2])) if len(clean) >= 2 else ''
    if combo:
        keys.insert(0, combo)
    return keys


def _row_group_marker(texts: list[str]) -> str:
    """Return the local subgroup used to disambiguate duplicate row labels."""
    joined = normalize_text(" ".join(texts))
    if "水平1" in joined:
        return "LEVEL1"
    if "水平2" in joined:
        return "LEVEL2"
    # Check outer before inner only for clarity; normalized Chinese does not overlap.
    if "外圈" in joined:
        return "OUTER"
    if "内圈" in joined:
        return "INNER"
    return ""


def _source_group_map(excel_rows: list[list[SourceCell]]) -> list[str]:
    groups: list[str] = []
    current = ""
    for row in excel_rows:
        marker = _row_group_marker(_source_row_texts(row))
        if marker:
            current = marker
        groups.append(current)
    return groups


def _labels_equivalent(left: str, right: str) -> bool:
    a = _canonical_row_label(left)
    b = _canonical_row_label(right)
    if not a or not b:
        return False
    if a == b:
        return True
    return len(a) >= 4 and len(b) >= 4 and (a in b or b in a)


def _labels_match_as_row_key(word_label: str, source_label: str) -> bool:
    """Match a row label without treating a narrative title as a data row."""
    if not _labels_equivalent(word_label, source_label):
        return False
    word_key = _canonical_row_label(word_label)
    source_key = _canonical_row_label(source_label)
    return word_key == source_key or normalize_text(source_label).startswith(normalize_text(word_label))


def _find_source_row(
    word_row_texts: list[str],
    excel_rows: list[list[SourceCell]],
    required_group: str = "",
    source_groups: list[str] | None = None,
) -> tuple[Optional[list[SourceCell]], int]:
    word_nonblank = [t for t in word_row_texts if not is_blank_value(t) and not _is_placeholder_text(t)]
    if not word_nonblank:
        return None, -1
    keys = _row_key_candidates(word_nonblank)
    if not keys:
        return None, -1
    groups = source_groups or [""] * len(excel_rows)

    def group_ok(index: int) -> bool:
        return not required_group or index >= len(groups) or groups[index] == required_group

    # 第一优先级：用 Word 当前行第一个非空内容作为行键。
    primary = word_nonblank[0]
    primary_norm = normalize_text(primary)
    for idx, srow in enumerate(excel_rows):
        if not group_ok(idx):
            continue
        stexts = _source_row_texts(srow)
        if len(stexts) < 1:
            continue
        first_text = stexts[0]
        first_norm = normalize_text(first_text)
        if _labels_match_as_row_key(primary, first_text):
            if primary_norm.isdigit() and _looks_like_numeric_header_source_row(srow):
                continue
            return srow, idx
        for root in ['最大吸光度', '相对偏倚', '准确性计算', '重复性计算', '稳定性计算', '结论', '指标', '波动百分比']:
            if root in primary_norm and first_norm.startswith(root):
                return srow, idx

    # 第二优先级：兼容行键不在 Excel 第一列，以及 Word 左侧存在纵向合并组名。
    for idx, srow in enumerate(excel_rows):
        if not group_ok(idx) or len(srow) < 1:
            continue
        stexts = _source_row_texts(srow)
        snorms = [normalize_text(t) for t in stexts]
        scombo2 = normalize_text(''.join(stexts[:2])) if len(stexts) >= 2 else ''
        for key in keys:
            if not key:
                continue
            if key == scombo2:
                return srow, idx
            if key.isdigit() and key in snorms and not _looks_like_numeric_header_source_row(srow):
                return srow, idx
            for source_text, source_norm in zip(stexts, snorms):
                if _labels_match_as_row_key(key, source_text):
                    return srow, idx
                for root in ['最大吸光度', '相对偏倚', '准确性计算', '重复性计算', '稳定性计算', '结论', '指标', '波动百分比']:
                    if root in key and source_norm.startswith(root):
                        return srow, idx
    return None, -1


def _source_data_after_key(word_row_texts: list[str], source_row: list[SourceCell], prefer_first_key: bool = False) -> list[SourceCell]:
    """根据 Word 行键定位 Excel 行中的数据起点。"""
    word_nonblank = [t for t in word_row_texts if not is_blank_value(t) and not _is_placeholder_text(t)]
    if not word_nonblank:
        return []
    stexts = [c.text for c in source_row]
    snorms = [normalize_text(t) for t in stexts]

    # 默认匹配 Word 行最后一个非空标签，例如线性表已有靶值 100 时从 100 右侧开始填；
    # 覆盖模式下则从第一行键右侧开始，便于覆盖模板中已有占位值/旧值。
    # 覆盖模式通常从第一行键右侧写，但遇到“测试数据/重复性计算 + 次数或Mean”等
    # 合并标签行时，真正行键是最后一个非空标签，必须从它右侧取数据，避免漏填前半段。
    if prefer_first_key and normalize_text(word_nonblank[0]) in HIGH_LEVEL_ROW_LABELS and len(word_nonblank) >= 2:
        key_text = word_nonblank[-1]
    else:
        key_text = word_nonblank[0] if prefer_first_key else word_nonblank[-1]
    key_norm = normalize_text(key_text)
    if key_norm in snorms:
        pos = snorms.index(key_norm)
        return source_row[pos + 1:]
    # Excel 与 Word 可能使用等价业务名称，例如“参考色素最大值/均值”、
    # “ALT(ALTn)/ALT”。按别名找到真实标签位置后再取右侧数据，
    # 避免把标签本身误写入 Word。
    for pos, source_text in enumerate(stexts):
        if _labels_equivalent(key_text, source_text):
            return source_row[pos + 1:]
    # 尝试前两个组合，例如 稳定性计算MAX
    if len(word_nonblank) >= 2:
        combo = normalize_text(''.join(word_nonblank[:2]))
        scombo = normalize_text(''.join(stexts[:2])) if len(stexts) >= 2 else ''
        if combo == scombo:
            return source_row[2:]
    # 退化：如果第一列是固定合并标签“测试数据”，Word 第二列是次数，Excel 第一列是次数
    if len(word_nonblank) >= 2 and normalize_text(word_nonblank[0]) in {'测试数据', '准确性计算', '重复性计算', '稳定性计算'}:
        key = normalize_text(word_nonblank[1])
        if key in snorms:
            return source_row[snorms.index(key) + 1:]
    # 再退化：跳过与 Word 非空标签数量相同的前缀
    return source_row[min(len(word_nonblank), len(source_row)):]


def _word_empty_tcs_after_key(table, row_index: int, overwrite_nonblank: bool) -> list[tuple[int, object]]:
    """返回 Word 当前行的数据单元格。

    旧逻辑从“第一个可写空单元格”开始填，遇到纵向合并标题导致的物理空格时，
    会把第一个数据误填到左侧标签列，后面的“反应盘外圈吸光度”等列就空了。
    现在改成：先找行键所在列（如 1、2、MAX、Mean、结论），数据从行键右侧开始填。
    """
    tcs = row_physical_tcs(table, row_index)
    texts = [tc_text(tc) for tc in tcs]
    nonblank = [(i, t.strip()) for i, t in enumerate(texts) if t and t.strip() and not _is_placeholder_text(t)]
    if not nonblank:
        return []

    # 非覆盖模式：从最后一个非占位标签右侧开始，避免写到左侧纵向合并标签列。
    # 覆盖模式：从第一个非占位行键右侧开始，确保已有旧值/占位值（如 PASS、携带污染率）也能被新 Excel 数据覆盖。
    if overwrite_nonblank and normalize_text(nonblank[0][1]) in HIGH_LEVEL_ROW_LABELS and len(nonblank) >= 2:
        key_col = nonblank[-1][0]
    else:
        key_col = nonblank[0][0] if overwrite_nonblank else nonblank[-1][0]
    start = key_col + 1

    if start >= len(tcs):
        return []
    return [(ci, tc) for ci, tc in enumerate(tcs[start:], start=start) if tc_is_writable(tc, overwrite_nonblank)]

def _looks_like_header_row(texts: list[str]) -> bool:
    joined = ' '.join(t for t in texts if t)
    keywords = [
        '次数', '采光周期', '检测项目', '测试项目', '校准位置参数', '要求',
        '吸光度', '结果', '温度值', '单位', '质控靶值', '范围低限',
        '范围高限', '实测值', '是否在控',
    ]
    if sum(1 for k in keywords if k in joined) >= 2 and not any(t.strip().isdigit() for t in texts if t.strip()):
        return True
    if '测试描述' in joined and sum(1 for t in texts if str(t).strip().startswith('第') and str(t).strip().endswith('组')) >= 2:
        return True
    normalized_cells = {normalize_text(t) for t in texts if t and str(t).strip()}
    if any('项目交叉污染' in normalize_text(t) for t in texts if t) and {'NA', 'K', 'CL'} & normalized_cells:
        return True
    first_norm = normalize_text(texts[0]) if texts else ''
    if first_norm in {'项目', '项目次数', '项目测试时间'} and {'NA', 'K', 'CL'} & normalized_cells:
        return True
    # 线性范围表的第二层表头：前两列空白，后面是 1/2/3，不能当作数据行。
    nonblank = [t.strip() for t in texts if t and t.strip()]
    if len(texts) >= 5 and (not texts[0].strip()) and (not texts[1].strip()) and nonblank and all(x.isdigit() for x in nonblank):
        return True
    return False




def _font_size_for_word_table(word: WordTableBlock) -> float:
    """在不改变表格布局的前提下尽量提高字号。

    v2 字号过小；v4 按列数自适应提高字号，同时仍固定表格布局，不改变列宽。
    """
    cols = max((len(r) for r in physical_table_matrix(word.table)), default=4)
    if cols <= 4:
        return 10.5
    if cols <= 8:
        return 10.0
    if cols <= 10:
        return 9.5
    if cols <= 12:
        return 9.0
    return 8.5



def _excel_block_has_raw_measurement_data(excel: ExcelBlock) -> bool:
    """Return True only when a source block has real measured inputs, not just targets/formulas."""
    object_key = normalize_text(excel.object_name + excel.title)
    for row in excel.rows:
        if not row:
            continue
        if _looks_like_numeric_header_source_row(row):
            continue
        label = normalize_text(row[0].text)
        raw_values = [cell for cell in row[1:] if not cell.is_formula and not is_blank_value(cell.value)]
        if not raw_values:
            continue

        # Linear tables always contain X/target values. They are not enough to
        # calculate averages, deviations, R/R² or conclusions without measured values.
        if "线性范围验证" in object_key and label.isdigit():
            if any(cell.col >= 5 for cell in raw_values):
                return True
            continue

        if any(token in label for token in (
            "电流", "吸光度", "实测", "测定值", "测量值",
            "样本针携带污染",
        )):
            return True
        if label.isdigit() or label in {"M1", "M2", "M3", "0H", "4H", "8H"} or label.startswith("实测"):
            return True
    return False


def _clear_cell_for_no_source(
    word: WordTableBlock,
    module: str,
    logs: list[FillLog],
    row_index: int,
    col_index: int,
    tc,
    font_size: float,
    reason: str,
) -> None:
    if not tc_text(tc).strip():
        return
    set_tc_text_keep_style(tc, "", font_size_pt=font_size)
    logs.append(FillLog(
        word_table_index=word.index,
        word_section=word.section,
        word_object=word.object_name,
        word_row=row_index + 1,
        word_col=col_index + 1,
        module=module,
        excel_sheet="",
        excel_section="",
        excel_object="",
        excel_row=0,
        excel_col=0,
        value="",
        reason=reason,
    ))


def _clear_empty_source_derived_values(word: WordTableBlock, excel: ExcelBlock, font_size: float) -> list[FillLog]:
    """Clear formula/template-derived values when the matched source has no measurements."""
    if _excel_block_has_raw_measurement_data(excel):
        return []

    logs: list[FillLog] = []
    derived_labels = {
        "结论", "B1%", "B2%", "B3%", "MEAN", "SD", "CV", "XMAX", "XMIN",
        "波动百分比", "携带污染率CLH", "携带污染率CHL",
    }
    for r_idx, row in enumerate(physical_table_matrix(word.table)):
        if not row:
            continue
        norms = [normalize_text(cell) for cell in row]
        joined = normalize_text(" ".join(row))
        tcs = row_physical_tcs(word.table, r_idx)
        reason = "源块无实测数据，清理模板/公式派生结果"

        if "统计" in joined or "验证结论" in joined:
            for col_idx, tc in enumerate(tcs):
                _clear_cell_for_no_source(word, excel.module, logs, r_idx, col_idx, tc, font_size, reason)
            continue

        first = norms[0]
        if first in derived_labels:
            for col_idx in range(1, len(tcs)):
                _clear_cell_for_no_source(word, excel.module, logs, r_idx, col_idx, tcs[col_idx], font_size, reason)
            continue

        # R²/R/a/b rows in the template use bare "R/a/b" placeholders as value cells.
        if {"R", "A", "B"} & set(norms):
            for col_idx, norm in enumerate(norms):
                if norm in {"R", "A", "B"} and col_idx < len(tcs):
                    _clear_cell_for_no_source(word, excel.module, logs, r_idx, col_idx, tcs[col_idx], font_size, reason)
    return logs


def _next_writable_after_label(table, row_index: int, label_col: int, next_label_col: int | None, overwrite_nonblank: bool):
    """在同一行内，找某个标签右侧、下一个标签左侧的第一个可写单元格。"""
    tcs = row_physical_tcs(table, row_index)
    end = next_label_col if next_label_col is not None else len(tcs)
    for ci in range(label_col + 1, min(end, len(tcs))):
        if tc_is_writable(tcs[ci], overwrite_nonblank):
            return ci, tcs[ci]
    return None


def _fill_label_value_pairs(word: WordTableBlock, excel: ExcelBlock, r_idx: int, row_texts: list[str], source_row: list[SourceCell], raw_only: bool, overwrite_nonblank: bool, font_size: float, allow_formula_values: bool) -> list[FillLog]:
    """处理一行多个“标签-值”对，例如 R²/R/a/b。普通按行键连续填充会只填最后一个标签后面的值。"""
    label_roots = ['R²', 'R', 'A', 'B']
    positions = [(i, t.strip()) for i, t in enumerate(row_texts) if t and t.strip() and not _is_placeholder_text(t)]
    norm_positions = [(i, normalize_text(t)) for i, t in positions]
    # 至少包含两个标签，才进入成对填充逻辑，避免干扰普通数据行。
    pair_labels = []
    for i, n in norm_positions:
        if n in {'R²', 'R', 'A', 'B'} or n.startswith('R²'):
            pair_labels.append((i, n))
    if len(pair_labels) < 2:
        return []

    snorms = [normalize_text(c.text) for c in source_row]
    logs: list[FillLog] = []
    for idx, (label_col, label_norm) in enumerate(pair_labels):
        next_col = pair_labels[idx + 1][0] if idx + 1 < len(pair_labels) else None
        target = _next_writable_after_label(word.table, r_idx, label_col, next_col, overwrite_nonblank)
        if not target:
            continue
        # 找 Excel 同名标签，取其右侧第一个有效值。
        src_label_idx = None
        for si, sn in enumerate(snorms):
            if sn == label_norm or (label_norm == 'R²' and sn.startswith('R²')):
                src_label_idx = si
                break
        if src_label_idx is None:
            continue
        src = None
        for sj in range(src_label_idx + 1, len(source_row)):
            cand = source_row[sj]
            if cand.is_formula and (raw_only or not allow_formula_values):
                continue
            if not is_blank_value(cand.value):
                src = cand
                break
        if src is None:
            continue
        target_col, tc = target
        set_tc_text_keep_style(tc, src.text, font_size_pt=font_size)
        logs.append(FillLog(
            word_table_index=word.index,
            word_section=word.section,
            word_object=word.object_name,
            word_row=r_idx + 1,
            word_col=target_col + 1,
            module=excel.module,
            excel_sheet=excel.sheet,
            excel_section=excel.section,
            excel_object=excel.object_name,
            excel_row=src.row,
            excel_col=src.col,
            value=src.text,
            reason='同一行多标签配对填充',
        ))
    return logs

def _canonical_header_name(text: str) -> str:
    value = normalize_text(text)
    if value in {"测试项目", "测定项目", "项目"}:
        return "ITEM"
    if "单位" in value and "次数" not in value:
        return "UNIT"
    if "质控靶值" in value or value == "靶值":
        return "TARGET"
    if "范围低限" in value or "下限" == value:
        return "LOW"
    if "范围高限" in value or "上限" == value:
        return "HIGH"
    if "实测值" in value or value == "测量值":
        return "MEASURED"
    if "是否在控" in value or "是否通过" in value or value == "结论":
        return "STATUS"
    return ""


def _word_accuracy_header_map(matrix: list[list[str]], row_index: int) -> dict[str, int]:
    for rr in range(row_index - 1, max(-1, row_index - 12), -1):
        mapping: dict[str, int] = {}
        for ci, text in enumerate(matrix[rr]):
            name = _canonical_header_name(text)
            if name:
                mapping[name] = ci
        if {"ITEM", "UNIT", "LOW", "HIGH", "MEASURED", "STATUS"}.issubset(mapping):
            return mapping
    return {}


def _excel_accuracy_header_map(excel: ExcelBlock, source_index: int) -> dict[str, int]:
    for rr in range(source_index - 1, max(-1, source_index - 12), -1):
        mapping: dict[str, int] = {}
        for cell in excel.rows[rr]:
            name = _canonical_header_name(cell.text)
            if name:
                mapping[name] = cell.col
        if {"ITEM", "UNIT", "LOW", "HIGH", "MEASURED", "STATUS"}.issubset(mapping):
            return mapping
    return {}


def _fill_clinical_accuracy_row(
    word: WordTableBlock,
    excel: ExcelBlock,
    matrix: list[list[str]],
    row_index: int,
    source_row: list[SourceCell],
    source_index: int,
    raw_only: bool,
    overwrite_nonblank: bool,
    font_size: float,
) -> list[FillLog]:
    word_map = _word_accuracy_header_map(matrix, row_index)
    excel_map = _excel_accuracy_header_map(excel, source_index)
    if not word_map or not excel_map:
        return []
    source_by_col = {cell.col: cell for cell in source_row}

    # 临床准确性必须有真实的“质控靶值”来源。
    # 当前 BS-5000 示例 Excel 中靶值为空，而低限/高限/是否在控是依赖靶值的公式。
    # 旧逻辑会把公式缓存中的 0/Fail 当成有效数据写入 Word，造成“无数据却自动生成”。
    # 这里严格执行：靶值为空时整行不填，单位、实测值和公式结果也不拼凑。
    target_source_col = excel_map.get("TARGET")
    target_source = source_by_col.get(target_source_col) if target_source_col is not None else None
    if target_source is None or is_blank_value(target_source.value):
        return []

    tcs = row_physical_tcs(word.table, row_index)
    logs: list[FillLog] = []
    for name in ("UNIT", "TARGET", "LOW", "HIGH", "MEASURED", "STATUS"):
        target_col = word_map.get(name)
        source_col = excel_map.get(name)
        if target_col is None or source_col is None or target_col >= len(tcs):
            continue
        source = source_by_col.get(source_col)
        if source is None or is_blank_value(source.value):
            continue
        if raw_only and source.is_formula:
            continue
        tc = tcs[target_col]
        if not tc_is_writable(tc, overwrite_nonblank):
            continue
        set_tc_text_keep_style(tc, source.text, font_size_pt=font_size)
        logs.append(FillLog(
            word_table_index=word.index,
            word_section=word.section,
            word_object=word.object_name,
            word_row=row_index + 1,
            word_col=target_col + 1,
            module=excel.module,
            excel_sheet=excel.sheet,
            excel_section=excel.section,
            excel_object=excel.object_name,
            excel_row=source.row,
            excel_col=source.col,
            value=source.text,
            reason='临床准确性按表头字段精确填充',
        ))
    return logs


def _fill_table_from_block(word: WordTableBlock, excel: ExcelBlock, raw_only: bool, overwrite_nonblank: bool) -> list[FillLog]:
    logs: list[FillLog] = []
    matrix = physical_table_matrix(word.table)
    font_size = _font_size_for_word_table(word)
    source_groups = _source_group_map(excel.rows)
    has_raw_measurements = _excel_block_has_raw_measurement_data(excel)
    current_word_group = ""
    for r_idx, row_texts in enumerate(matrix):
        marker = _row_group_marker(row_texts)
        if marker:
            current_word_group = marker
        if _looks_like_header_row(row_texts):
            continue
        source_row, src_index = _find_source_row(
            row_texts,
            excel.rows,
            required_group=current_word_group,
            source_groups=source_groups,
        )
        if source_row is None:
            continue
        if _word_block_kind(word) == "clinical_accuracy" and _excel_block_kind(excel) == "clinical_accuracy":
            accuracy_logs = _fill_clinical_accuracy_row(
                word, excel, matrix, r_idx, source_row, src_index, raw_only, overwrite_nonblank, font_size
            )
            logs.extend(accuracy_logs)
            # 临床准确性由专用表头映射独占处理。即使因靶值缺失而没有写入，
            # 也不能再退回普通连续填充，否则仍会把单位/0/Fail 错填进去。
            continue
        # 先处理 R²/R/a/b 这类一行多个标签-值对。
        pair_logs = _fill_label_value_pairs(
            word, excel, r_idx, row_texts, source_row, raw_only, overwrite_nonblank, font_size, has_raw_measurements
        )
        if pair_logs:
            logs.extend(pair_logs)
            continue
        target_cells = _word_empty_tcs_after_key(word.table, r_idx, overwrite_nonblank)
        if not target_cells:
            continue
        src_values = _source_data_after_key(row_texts, source_row, prefer_first_key=overwrite_nonblank)
        if not src_values:
            continue
        fill_i = 0
        for target_col, tc in target_cells:
            # 跳过公式源单元格：Word 有公式时只填原始数据；Word 无公式时才同步 Excel 结果
            while fill_i < len(src_values) and (src_values[fill_i].is_formula and (raw_only or not has_raw_measurements)):
                fill_i += 1
            if fill_i >= len(src_values):
                break
            src = src_values[fill_i]
            fill_i += 1
            if is_blank_value(src.value):
                continue
            set_tc_text_keep_style(tc, src.text, font_size_pt=font_size)
            logs.append(FillLog(
                word_table_index=word.index,
                word_section=word.section,
                word_object=word.object_name,
                word_row=r_idx + 1,
                word_col=target_col + 1,
                module=excel.module,
                excel_sheet=excel.sheet,
                excel_section=excel.section,
                excel_object=excel.object_name,
                excel_row=src.row,
                excel_col=src.col,
                value=src.text,
                reason='章节+名称+表头匹配后按行键复制',
            ))
    logs.extend(_clear_empty_source_derived_values(word, excel, font_size))
    return logs


def _clear_unmatched_template_values(word: WordTableBlock, module: str) -> list[FillLog]:
    """Clear stale template result text for tables that have no trusted Excel source."""
    if _word_block_kind(word) != "gravimetric_sampling":
        return []

    logs: list[FillLog] = []
    font_size = _font_size_for_word_table(word)
    matrix = physical_table_matrix(word.table)
    for r_idx, row in enumerate(matrix):
        first = normalize_text(row[0]) if row else ""
        second = normalize_text(row[1]) if len(row) > 1 else ""
        joined = normalize_text(" ".join(row))
        if first != "结论" and not (second in {"正确度", "重复性"} and "符合要求" in joined):
            continue

        tcs = row_physical_tcs(word.table, r_idx)
        # Preserve the row labels ("结论/正确度/重复性") but remove template verdicts.
        for col_idx in range(2, len(tcs)):
            old_text = tc_text(tcs[col_idx]).strip()
            if not old_text:
                continue
            set_tc_text_keep_style(tcs[col_idx], "", font_size_pt=font_size)
            logs.append(FillLog(
                word_table_index=word.index,
                word_section=word.section,
                word_object=word.object_name,
                word_row=r_idx + 1,
                word_col=col_idx + 1,
                module=module,
                excel_sheet="",
                excel_section="",
                excel_object="",
                excel_row=0,
                excel_col=0,
                value="",
                reason="无可信数据源，清理模板预置结果",
            ))
    return logs


def _infer_module_for_table(word: WordTableBlock, occurrence_by_key: dict[str, int], has_m2: bool) -> str:
    if getattr(word, 'module_hint', ''):
        return word.module_hint
    text = normalize_text(word.key_text)
    if 'M2' in text and 'M1' not in text:
        return 'M2'
    if 'M1' in text and 'M2' not in text:
        return 'M1'
    if not has_m2:
        return 'M1'
    # 如果模板/标准答案里同一表连续出现两次，第一次按 M1、第二次按 M2
    key = normalize_text(word.section + '|' + word.title + '|' + word.object_name)
    count = occurrence_by_key.get(key, 0)
    occurrence_by_key[key] = count + 1
    return 'M1' if count % 2 == 0 else 'M2'



def _delete_top_sections_from(doc, start_nums: set[str] = {'六', '七', '八', '九'}) -> int:
    """删除从“六、...”开始到文档结尾的所有正文块。

    用户要求六、七、八、九章节不要生成；这里删除顶层章节，不影响（六）（七）这类检测子章节。
    """
    body = doc.element.body
    children = list(body)
    delete_from = None
    for idx, child in enumerate(children):
        texts = []
        for t in child.iter():
            if t.tag.endswith('}t') and t.text:
                texts.append(t.text)
        text = ''.join(texts).strip()
        if not text:
            continue
        if looks_like_top_section(text):
            first = text[0]
            if first in start_nums:
                delete_from = idx
                break
    if delete_from is None:
        return 0
    removed = 0
    for child in children[delete_from:]:
        try:
            body.remove(child)
            removed += 1
        except Exception:
            pass
    return removed


def _prepare_tables_for_fixed_fill(doc) -> None:
    for table in doc.tables:
        set_table_fixed_layout(table)

def fill_report_tables(
    template_path: str | Path,
    output_path: str | Path,
    m1_excel: str | Path,
    m2_excel: str | Path | None = None,
    formula_policy: str = 'all',
    overwrite_nonblank: bool = True,
    delete_tail_sections: bool = True,
) -> tuple[list[FillLog], list[MatchLog], list[str]]:
    """按“模块 -> 章节 -> 名称 -> 表头 -> 行键”填充 Word 表格。

    formula_policy:
      - auto：若 Word 模板有表格公式，只填 Excel 原始数据；若没有公式，同步 Excel 已计算结果。
      - raw：只填 Excel 原始数据，跳过 Excel 公式单元格。
      - all：连 Excel 公式的缓存结果也填入 Word。
    """
    warnings: list[str] = []
    formulas = find_word_table_formulas(template_path)
    if formula_policy == 'auto':
        raw_only = bool(formulas)
        if formulas:
            warnings.append(f'检测到 Word 表格公式 {len(formulas)} 个：仅复制 Excel 原始数据，公式结果由 Word 更新。')
        else:
            raw_only = False
            warnings.append('未检测到 Word 表格计算公式：已同步 Excel 中的计算结果，避免报告结果为空。')
    elif formula_policy == 'raw':
        raw_only = True
    else:
        raw_only = False

    excel_blocks = read_excel_blocks(m1_excel, 'M1')
    if m2_excel:
        excel_blocks.extend(read_excel_blocks(m2_excel, 'M2'))

    doc = Document(str(template_path))
    removed_tail_blocks = _delete_top_sections_from(doc) if delete_tail_sections else 0
    if removed_tail_blocks:
        warnings.append(f'已删除顶层六、七、八、九章节相关正文块：{removed_tail_blocks} 个。')
    _prepare_tables_for_fixed_fill(doc)
    word_blocks = get_word_table_blocks(doc)
    has_m2 = any(block.module == 'M2' for block in excel_blocks)
    occurrence_by_key: dict[str, int] = {}
    fill_logs: list[FillLog] = []
    match_logs: list[MatchLog] = []

    for wb in word_blocks:
        module = _infer_module_for_table(wb, occurrence_by_key, has_m2)

        # BS-5000“暗电流检测”在 Word 中是一个物理表格，但 Excel 分成
        # “内圈暗电流测试”和“外圈暗电流测试”两个数据块。这里局部、精确地
        # 分别填充两个区域，禁止修改通用评分/行匹配规则，避免修复一个章节后
        # 影响其他章节。
        split_sources = _dark_current_split_sources(wb, excel_blocks, module)
        if split_sources is not None:
            inner_src, outer_src = split_sources
            if inner_src is None or outer_src is None:
                missing_parts = []
                if inner_src is None:
                    missing_parts.append('内圈')
                if outer_src is None:
                    missing_parts.append('外圈')
                match_logs.append(MatchLog(
                    word_table_index=wb.index, word_section=wb.section, word_object=wb.object_name,
                    module=module, excel_sheet='', excel_start_row=0, excel_end_row=0,
                    excel_section='暗电流检测', excel_object='', score=0.0,
                    note='暗电流精确匹配缺少' + '/'.join(missing_parts) + '数据块，未填充',
                ))
                continue

            match_logs.append(MatchLog(
                word_table_index=wb.index, word_section=wb.section, word_object='内圈+外圈暗电流测试',
                module=module, excel_sheet=inner_src.sheet,
                excel_start_row=min(inner_src.start_row, outer_src.start_row),
                excel_end_row=max(inner_src.end_row, outer_src.end_row),
                excel_section='暗电流检测', excel_object='内圈暗电流测试 + 外圈暗电流测试',
                score=200.0, note='已匹配',
            ))
            fill_logs.extend(_fill_table_from_block(
                wb, inner_src, raw_only=raw_only, overwrite_nonblank=overwrite_nonblank
            ))
            fill_logs.extend(_fill_table_from_block(
                wb, outer_src, raw_only=raw_only, overwrite_nonblank=overwrite_nonblank
            ))
            continue

        src, score = _choose_block(wb, excel_blocks, module)
        if src is None:
            fill_logs.extend(_clear_unmatched_template_values(wb, module))
            match_logs.append(MatchLog(
                word_table_index=wb.index, word_section=wb.section, word_object=wb.object_name,
                module=module, excel_sheet='', excel_start_row=0, excel_end_row=0,
                excel_section='', excel_object='', score=score, note='未找到可信匹配，未填充',
            ))
            continue
        match_logs.append(MatchLog(
            word_table_index=wb.index, word_section=wb.section, word_object=wb.object_name,
            module=module, excel_sheet=src.sheet, excel_start_row=src.start_row, excel_end_row=src.end_row,
            excel_section=src.section, excel_object=src.object_name, score=score, note='已匹配',
        ))
        fill_logs.extend(_fill_table_from_block(wb, src, raw_only=raw_only, overwrite_nonblank=overwrite_nonblank))

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return fill_logs, match_logs, warnings


def save_logs(fill_logs: list[FillLog], match_logs: list[MatchLog], out_prefix: str | Path) -> tuple[Path, Path]:
    out_prefix = Path(out_prefix)
    fill_path = out_prefix.with_name(out_prefix.stem + '_填充日志.csv')
    match_path = out_prefix.with_name(out_prefix.stem + '_匹配日志.csv')
    fill_path.parent.mkdir(parents=True, exist_ok=True)
    with fill_path.open('w', encoding='utf-8-sig', newline='') as f:
        fields = list(FillLog.__dataclass_fields__.keys())
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for item in fill_logs:
            w.writerow(asdict(item))
    with match_path.open('w', encoding='utf-8-sig', newline='') as f:
        fields = list(MatchLog.__dataclass_fields__.keys())
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for item in match_logs:
            w.writerow(asdict(item))
    return fill_path, match_path
